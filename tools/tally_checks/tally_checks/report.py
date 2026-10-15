"""Report writer: report.xlsx, report.json, summary.md, delete_list.csv."""
import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

SEV_ORDER = {"high": 0, "medium": 1, "low": 2}
FIXED_TS = datetime(2000, 1, 1)


def _sheet(wb, title, rows):
    ws = wb.create_sheet(title[:31])
    if not rows:
        ws.append(["No findings"])
        return
    cols = list(rows[0].keys())
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r[c] for c in cols])
    for i, c in enumerate(cols, 1):
        width = max(len(str(c)), *(len(str(r[c])) for r in rows)) + 2
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(width, 60)


def write_report(results, meta, out_dir):
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    results = sorted(results, key=lambda r: r.id)

    # JSON (deterministic: no timestamps, sorted order)
    doc = {"meta": meta, "checks": [{
        "id": r.id, "name": r.name, "severity": r.severity, "count": r.count,
        "total_amount": r.total_amount, "findings": r.rows, "detail": r.detail,
    } for r in results]}
    (out / "report.json").write_text(json.dumps(doc, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    # XLSX
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.created = wb.properties.modified = FIXED_TS
    _sheet(wb, "Summary", [{"check": f"{r.id}. {r.name}", "count": r.count,
                            "total_amount": r.total_amount, "severity": r.severity} for r in results])
    for r in results:
        _sheet(wb, f"{r.id}_{r.name}", r.rows)
        if r.detail:
            _sheet(wb, f"{r.id}_detail", r.detail)
    wb.save(out / "report.xlsx")

    # delete_list.csv (checks 1 and 2 only; never delete XML)
    cands = {}
    for r in results:
        if r.id in (1, 2):
            for d in r.delete_candidates:
                cands[(d["voucher_no"], d["date"], d["voucher_type"])] = d
    with open(out / "delete_list.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["voucher_no", "date", "voucher_type", "amount"])
        for k in sorted(cands, key=lambda k: (k[1], k[0], k[2])):
            d = cands[k]
            w.writerow([d["voucher_no"], d["date"], d["voucher_type"], f"{d['amount']:.2f}"])

    # summary.md
    lines = [f"# Tally checks: {meta['company']}",
             f"Period {meta['from'] or 'start'} to {meta['to'] or 'end'} · "
             f"{meta['voucher_lines']} voucher lines · source {meta['daybook']}", ""]
    for r in sorted(results, key=lambda r: (SEV_ORDER.get(r.severity, 9), r.id)):
        flag = "OK  " if r.count == 0 else r.severity.upper()
        lines.append(f"- [{flag}] #{r.id} {r.name}: {r.count} finding(s), total {r.total_amount:,.2f}")
    if cands:
        lines += ["", f"{len(cands)} suspected phantom voucher(s) listed in delete_list.csv "
                      "for manual review and deletion in Tally."]
    lines += ["", "Scripts are read-only. Nothing was written to Tally."]
    text = "\n".join(lines) + "\n"
    (out / "summary.md").write_text(text, encoding="utf-8")
    return text
